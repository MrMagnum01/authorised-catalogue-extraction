import csv
from pathlib import Path

import httpx
from openpyxl import load_workbook

from catalogue_extract.boundary import Scope
from catalogue_extract.crawl import CrawlConfig, run_crawl
from catalogue_extract.export import sanitize_formula_field, write_accounting_report, write_products_csv, write_products_xlsx
from catalogue_extract.snapshot import crawl_result_to_dict, publish, load_generation
from catalogue_fixture.data import FixtureProduct
from catalogue_fixture.site import FixtureServer, FixtureState


def test_sanitize_formula_field_prefixes_dangerous_leading_chars():
    assert sanitize_formula_field("=1+1") == "'=1+1"
    assert sanitize_formula_field("+SUM(A1)") == "'+SUM(A1)"
    assert sanitize_formula_field("-1") == "'-1"
    assert sanitize_formula_field("@cmd") == "'@cmd"
    assert sanitize_formula_field("Normal Name") == "Normal Name"
    assert sanitize_formula_field(None) is None


def _client():
    return httpx.Client(trust_env=False)


def test_csv_and_xlsx_export_with_injection_attempt(tmp_path: Path):
    dangerous = FixtureProduct("SKU-EVIL", "=cmd|'/c calc'!A1", 999, "USD", "Widgets")
    state = FixtureState(products=[dangerous])
    with FixtureServer(state) as server:
        scope = Scope.from_origin(server.origin, ["/catalogue"])
        with _client() as client:
            result = run_crawl(
                client, scope,
                seed_listing_url=server.origin + "/catalogue/",
                permission_url=server.origin + "/PERMISSION.md",
                robots_url=server.origin + "/robots.txt",
                snapshot_meta_url=server.origin + "/_meta/snapshot",
                config=CrawlConfig(min_interval_s=0.02, jitter_s=0.0),
            )
        publish(tmp_path, result)

    generation = load_generation(tmp_path, result.crawl_id)

    # The raw generation JSON keeps the untransformed value.
    raw_name = generation["products"][0]["name"]
    assert raw_name == "=cmd|'/c calc'!A1"

    csv_path = tmp_path / "products.csv"
    write_products_csv(generation, csv_path)
    with open(csv_path, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    assert rows[0]["name"].startswith("'=")
    assert rows[0]["product_id"] == "SKU-EVIL"
    assert rows[0]["content_sha256"]
    assert rows[0]["parser_version"]

    xlsx_path = tmp_path / "products.xlsx"
    write_products_xlsx(generation, xlsx_path)
    wb = load_workbook(xlsx_path)
    ws = wb.active
    header = [c.value for c in ws[1]]
    name_col = header.index("name")
    assert ws.cell(row=2, column=name_col + 1).value.startswith("'=")

    report_path = tmp_path / "report.json"
    write_accounting_report(generation, report_path)
    assert report_path.exists()


def test_provenance_fields_present_per_row(tmp_path: Path):
    state = FixtureState(products=[FixtureProduct("SKU-P", "Plain", 100, "USD", "Widgets")])
    with FixtureServer(state) as server:
        scope = Scope.from_origin(server.origin, ["/catalogue"])
        with _client() as client:
            result = run_crawl(
                client, scope,
                seed_listing_url=server.origin + "/catalogue/",
                permission_url=server.origin + "/PERMISSION.md",
                robots_url=server.origin + "/robots.txt",
                snapshot_meta_url=server.origin + "/_meta/snapshot",
                config=CrawlConfig(min_interval_s=0.02, jitter_s=0.0),
            )
    generation = crawl_result_to_dict(result)
    csv_path = tmp_path / "products.csv"
    write_products_csv(generation, csv_path)
    with open(csv_path, newline="", encoding="utf-8") as f:
        row = next(csv.DictReader(f))
    assert row["source_url"]
    assert row["fetch_utc"]
    assert row["http_status"] == "200"
    assert row["content_sha256"]
    assert row["parser_version"] == "1.0.0"
